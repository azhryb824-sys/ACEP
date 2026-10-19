"""Extract original project BOQ/actuals with header-driven mapping, no fitting."""
from pathlib import Path
import json, hashlib, numbers, math, datetime
import openpyxl
import pandas as pd

R=Path(__file__).resolve().parent;O=R/'boq_data';O.mkdir(exist_ok=True)
def number(v):return isinstance(v,numbers.Number) and not isinstance(v,bool) and math.isfinite(float(v))
def norm(v):return str(v or '').strip().lower()
aliases={'hr':'h','hrs':'h','hour':'h','hours':'h','m2':'m2','sqm':'m2','m²':'m2','m3':'m3','cum':'m3','m³':'m3','lm':'m','m':'m','item':'item','ls':'item','no':'each','no.':'each','ea':'each','each':'each','tonne':'t','tonnes':'t','t':'t','ton':'t','weeks':'week','week':'week','days':'day','day':'day','month':'month','night':'night','nights':'night','kg':'kg','km':'km','ha':'ha'}
records=[];reject=[];projects=[];resources=[];actual_entries=[];networks=[]
for p in sorted((R/'sources').glob('P*.xlsx')):
 pid=p.name.split('_')[0];w=openpyxl.load_workbook(p,read_only=True,data_only=True)
 s=w['Portfolio WBS'];rr=list(s.iter_rows(values_only=True))
 hi=next(i for i,row in enumerate(rr[:5]) if any(norm(v)=='final qty' for v in row))
 head={norm(v):j for j,v in enumerate(rr[hi]) if isinstance(v,str) and norm(v)}
 q=next(head[x] for x in ('quantity','qty') if x in head)
 u=next(head[x] for x in ('uom','unit') if x in head)
 budget=head['original budget'] if 'original budget' in head else head.get('bac',head.get('budget'))
 assert budget is not None
 count=0
 for ri,row in enumerate(rr[hi+1:],hi+2):
  if not number(row[0]):continue
  fields=[row[q],row[budget],row[head['final qty']],row[head['final cost']]]
  if not all(number(x) for x in fields):reject.append({'project_id':pid,'row':ri,'reason':'missing_or_nonfinite_amount'});continue
  planned,bac,aq,ac=map(float,fields);original_unit=norm(row[u]);unit=aliases.get(original_unit)
  rec=dict(project_id=pid,item_id=f'{pid}-WBS-{row[0]:g}',wbs=str(row[0]),description=str(row[head['description']] or '').strip(),unit=unit,original_unit=original_unit,planned_quantity=planned,planned_cost_aud=bac,actual_quantity=aq,actual_cost_aud=ac,source_file=p.name,source_excel_row=ri,budget_header=rr[hi][budget])
  if not unit:reject.append(dict(rec,reason='unrecognized_or_missing_unit'));continue
  if min(planned,bac,aq,ac)<0:reject.append(dict(rec,reason='negative_amount'));continue
  records.append(rec);count+=1
 # Preserve explicit dependency strings without converting SS/FF/lag semantics to FS.
 s=w['Program Links']
 for ri,row in enumerate(s.iter_rows(min_row=2,values_only=True),2):
  if number(row[0]) and isinstance(row[1],str):networks.append(dict(project_id=pid,wbs=str(row[0]),description=row[1],duration=row[2] if number(row[2]) else None,predecessors=str(row[3] or '').strip(),successors=str(row[4] or '').strip(),source_excel_row=ri))
 # Actual costs are retained as observed transaction data, not silently summed with BOQ estimates.
 ah=next(w['Actual Costs'].iter_rows(max_row=1,values_only=True))
 ah={norm(v):j for j,v in reversed(list(enumerate(ah[:10]))) if isinstance(v,str) and norm(v)}
 total_col=ah['total'];wbs_col=ah.get('portfolio wbs',ah.get('cost code'))
 assert wbs_col is not None
 for ri,row in enumerate(w['Actual Costs'].iter_rows(min_row=2,values_only=True),2):
  if isinstance(row[0],datetime.datetime) and all(number(row[j]) for j in (ah['quantity'],ah['rate'],total_col)):
   qty=float(row[ah['quantity']]);rate=float(row[ah['rate']]);total=float(row[total_col]);adjustment=row[ah['adjustment']] if 'adjustment' in ah else None
   actual_entries.append(dict(project_id=pid,date=row[0].isoformat()[:10],resource=str(row[ah['resource']] or '').strip(),unit=str(row[ah['unit']] or '').strip(),quantity=qty,rate_aud=rate,total_aud=total,wbs=str(row[wbs_col] or '').strip(),source_excel_row=ri,adjustment=adjustment if number(adjustment) else None,arithmetic_residual_aud=total-qty*rate))
 projects.append(dict(project_id=pid,file=p.name,sha256=hashlib.sha256(p.read_bytes()).hexdigest(),boq_rows=count))
df=pd.DataFrame(records);assert not df.item_id.duplicated().any()
df.to_csv(O/'all_audited_boq.csv',index=False)
pd.DataFrame(actual_entries).to_csv(O/'actual_resource_transactions.csv',index=False)
pd.DataFrame(networks).to_csv(O/'activity_dependencies.csv',index=False)
(O/'quarantined_rows.json').write_text(json.dumps(reject,indent=2)+'\n')
# Eligibility depends only on the presence of a positive baseline and a recognized unit.
eligible=df.planned_quantity.gt(0)&df.planned_cost_aud.gt(0)
split={'P01':'train','P17':'train','P18':'validation','P02':'test','P03':'test','P04':'test','P05':'test','P06':'test'}
features=['item_id','project_id','wbs','description','unit','planned_quantity','planned_cost_aud']
labels=['item_id','actual_quantity','actual_cost_aud']
for name in ('train','validation','test'):
 part=df[eligible & df.project_id.map(split).eq(name)]
 part[features].to_csv(O/f'{name}_features.csv',index=False,float_format='%.12g')
 part[labels].to_csv(O/f'{name}_labels.csv',index=False,float_format='%.12g')
contract={'version':'5b.0','seed':42196,'source':'10.6084/m9.figshare.12998822.v3','license':'CC BY 4.0','train_projects':['P01','P17'],'validation_projects':['P18'],'test_projects':['P02','P03','P04','P05','P06'],'tasks':['actual_quantity','actual_cost_aud'],'eligibility':'Positive supplied baseline quantity and budget; recognized unit; finite nonnegative observed outcomes. Zero final quantities/costs are retained and counted as failures unless predicted as zero. New work with no initial baseline is outside this conditional predictor and reported in coverage.','selection':'Highest mean per-project fraction within 10% on validation; tie lowest mean per-item bounded relative error. No test-based tuning or refitting.','relative_error':'abs(pred-actual)/actual for actual>0; 0 if both zero, otherwise failure. Selection uses bounded error abs(pred-actual)/max(abs(pred),abs(actual),1e-12) to avoid infinities.','gate':'For each of five projects: at least 90% of all audited reference lines must have a covered prediction within 10%. Report eligible-only and all-audited-line coverage separately. Cost also requires aggregate relative error <=10%. Never sum unlike quantity units.','scope':'Conditional final BOQ quantity and cost from a supplied estimate, historical Australian civil works. Not quantity takeoff from drawings, not Saudi prices.','limitations':['These projects were used in earlier duration research. This is a separately held-out BOQ task, not five wholly unseen projects across all ACEP engines.','A few source rows were inspected for schema audit before training; outcomes were not used for candidate choice.','Original workbooks are retrospective and do not provide full budget revision history.','Small portfolio, one validation project, shared organization/resources; no external generalization proof.'],'candidates':['unchanged_baseline','global_median_ratio','unit_median_ratio','ridge_1','ridge_10','ridge_100','extra_2','extra_5','hist_7','neural_32_16']}
(O/'contract.json').write_text(json.dumps(contract,indent=2)+'\n')
audit={'projects':projects,'audited_boq_rows':len(df),'eligible_rows':int(eligible.sum()),'zero_baseline_rows':int((~eligible).sum()),'quarantined_rows':len(reject),'resource_transactions':len(actual_entries),'dependency_rows':len(networks),'counts':{n:int((eligible & df.project_id.map(split).eq(n)).sum()) for n in ('train','validation','test')},'per_project_counts':df.groupby('project_id').size().to_dict(),'per_project_eligible':df[eligible].groupby('project_id').size().to_dict(),'transaction_arithmetic_discrepancies_over_0_05':sum(abs(x['arithmetic_residual_aud'])>.05 for x in actual_entries),'note':'Resource transaction totals and final BOQ cached amounts are retained separately; discrepancies are not silently reconciled.'}
(O/'audit.json').write_text(json.dumps(audit,indent=2)+'\n')
(O/'manifest.json').write_text(json.dumps({p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(O.iterdir()) if p.name!='manifest.json'},indent=2)+'\n')
print(json.dumps(audit,indent=2))
